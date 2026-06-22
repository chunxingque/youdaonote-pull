#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将有道云笔记的 .lxtable 文件转换为 Markdown 表格或 Excel
用法：
    python convert_lxtable.py <input.lxtable> [output.md|output.xlsx]
"""

import json
import sys
import os
from typing import List, Dict, Any


def parse_lxtable(file_path: str) -> List[Dict[str, Any]]:
    """
    解析 .lxtable 文件，返回表格数据
    返回格式：[{"name": sheet_name, "headers": [header1, header2, ...], "rows": [[cell1, cell2, ...], ...]}, ...]
    """
    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    sheets_data = []
    sheets = data.get('data', {}).get('sheets', {})
    tabs = data.get('data', {}).get('tabs', ['0'])
    
    for sheet_idx, (sheet_key, sheet_data) in enumerate(sheets.items()):
        cells = sheet_data.get('cells', {})
        
        # 找出表格的最大行和列
        max_row, max_col = 0, 0
        for cell_key in cells.keys():
            try:
                row, col = map(int, cell_key.split(','))
                max_row = max(max_row, row)
                max_col = max(max_col, col)
            except:
                pass
        
        # 提取表头（假设第一行是表头）
        headers = []
        for col in range(max_col + 1):
            cell_key = f'0,{col}'
            cell_data = cells.get(cell_key, {})
            header_value = cell_data.get('0', f'列{col+1}')
            headers.append(str(header_value))
        
        # 提取数据行
        rows = []
        for row in range(1, max_row + 1):
            row_data = []
            for col in range(max_col + 1):
                cell_key = f'{row},{col}'
                cell_data = cells.get(cell_key, {})
                cell_value = cell_data.get('0', '')
                row_data.append(str(cell_value) if cell_value else '')
            rows.append(row_data)
        
        sheet_name = tabs[sheet_idx] if sheet_idx < len(tabs) else f'Sheet{sheet_idx+1}'
        sheets_data.append({
            'name': sheet_name,
            'headers': headers,
            'rows': rows
        })
    
    return sheets_data


def convert_to_markdown(sheets_data: List[Dict[str, Any]], title: str = '') -> str:
    """将表格数据转换为 Markdown 格式"""
    md_lines = []
    
    if title:
        md_lines.append(f'# {title}\n')
    
    for sheet in sheets_data:
        md_lines.append(f'## {sheet["name"]}\n')
        
        # 表头
        md_lines.append('| ' + ' | '.join(sheet['headers']) + ' |')
        md_lines.append('| ' + ' | '.join(['---'] * len(sheet['headers'])) + ' |')
        
        # 数据行
        for row in sheet['rows']:
            # 转义 Markdown 特殊字符
            escaped_row = [cell.replace('|', '\\|').replace('\n', '<br/>') for cell in row]
            md_lines.append('| ' + ' | '.join(escaped_row) + ' |')
        
        md_lines.append('')  # 空行
    
    return '\n'.join(md_lines)


def convert_to_excel(sheets_data: List[Dict[str, Any]], output_path: str, title: str = ''):
    """将表格数据转换为 Excel 文件"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except ImportError:
        print('错误：需要安装 openpyxl 库。请运行：pip install openpyxl')
        return False
    
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认 sheet
    
    for sheet in sheets_data:
        ws = wb.create_sheet(title=sheet['name'][:31])  # Excel sheet 名最多 31 字符
        
        # 写入表头
        ws.append(sheet['headers'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        # 写入数据行
        for row in sheet['rows']:
            ws.append(row)
    
    wb.save(output_path)
    return True


def convert_lxtable_file(input_path: str, output_path: str = None, format: str = 'md'):
    """
    转换 .lxtable 文件
    :param input_path: 输入的 .lxtable 文件路径
    :param output_path: 输出文件路径（可选，默认与输入同目录同名）
    :param format: 输出格式，'md' 或 'xlsx'
    """
    if not os.path.exists(input_path):
        print(f'错误：文件不存在：{input_path}')
        return False
    
    # 解析数据
    try:
        sheets_data = parse_lxtable(input_path)
    except Exception as e:
        print(f'错误：解析文件失败：{e}')
        return False
    
    # 确定输出路径
    if not output_path:
        base_name = os.path.splitext(input_path)[0]
        output_path = f'{base_name}.{format}'
    
    # 转换
    try:
        if format == 'md':
            title = os.path.basename(input_path).replace('.lxtable', '')
            md_content = convert_to_markdown(sheets_data, title)
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f'成功：已转换为 Markdown 表格：{output_path}')
            
        elif format == 'xlsx':
            title = os.path.basename(input_path).replace('.lxtable', '')
            if convert_to_excel(sheets_data, output_path, title):
                print(f'成功：已转换为 Excel 文件：{output_path}')
            else:
                return False
        else:
            print(f'错误：不支持的格式：{format}')
            return False
            
    except Exception as e:
        print(f'错误：转换失败：{e}')
        return False
    
    return True


def batch_convert(directory: str, format: str = 'md'):
    """批量转换目录中的所有 .lxtable 文件"""
    count = 0
    for root, dirs, files in os.walk(directory):
        for file in files:
            if file.endswith('.lxtable'):
                input_path = os.path.join(root, file)
                print(f'处理：{input_path}')
                if convert_lxtable_file(input_path, format=format):
                    count += 1
    
    print(f'\n完成：共转换 {count} 个文件')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法：')
        print('  转换单个文件：python convert_lxtable.py <input.lxtable> [output.md|output.xlsx]')
        print('  批量转换目录：python convert_lxtable.py --batch <directory> [--format md|xlsx]')
        sys.exit(1)
    
    if sys.argv[1] == '--batch':
        directory = sys.argv[2] if len(sys.argv) > 2 else '.'
        format = 'md'
        if '--format' in sys.argv:
            idx = sys.argv.index('--format')
            if idx + 1 < len(sys.argv):
                format = sys.argv[idx + 1]
        batch_convert(directory, format)
    else:
        input_path = sys.argv[1]
        output_path = sys.argv[2] if len(sys.argv) > 2 else None
        format = 'xlsx' if output_path and output_path.endswith('.xlsx') else 'md'
        convert_lxtable_file(input_path, output_path, format)
